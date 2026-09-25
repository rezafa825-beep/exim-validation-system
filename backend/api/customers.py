from fastapi import APIRouter, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from database.database import get_conn
import io
from openpyxl import Workbook, load_workbook

router=APIRouter(prefix="/api/customers",tags=["customers"])

class CustomerIn(BaseModel):
    customer_name:str
    status:str="ACTIVE"
    mappings:dict[str,dict[str,str|list[str]]] = Field(default_factory=dict)
    row_rule:dict = Field(default_factory=dict)

def norm(s):
    return " ".join((s or "").replace("PT.","").replace("PT ","").split()).upper()

def _split_headers(value):
    if isinstance(value,list): return [str(x).strip() for x in value if str(x).strip()]
    return [x.strip() for x in str(value or "").replace("\n",";").split(";") if x.strip()]

def _headers(primary, alternatives):
    return ";".join(_split_headers(alternatives))

def _mapping_rows(c, cid):
    return [dict(r) for r in c.execute("SELECT document_type,standard_field,primary_header,alternative_headers FROM customer_mappings WHERE customer_id=? ORDER BY document_type,standard_field",(cid,))]

@router.get("")
def customers():
    c=get_conn(); rows=[dict(r) for r in c.execute("SELECT * FROM customers ORDER BY customer_name")]; c.close(); return rows

@router.get("/id/{customer_id}")
def get_customer(customer_id:int):
    c=get_conn(); row=c.execute("SELECT * FROM customers WHERE id=?",(customer_id,)).fetchone()
    if not row: c.close(); raise HTTPException(404,"Customer tidak ditemukan")
    mappings={}
    for r in c.execute("SELECT document_type,standard_field,primary_header,alternative_headers FROM customer_mappings WHERE customer_id=?",(customer_id,)):
        mappings.setdefault(r[0],{})[r[1]]={"primary":r[2] or "","alternatives":_split_headers(r[3])}
    rr=c.execute("SELECT * FROM row_rules WHERE customer_id=? AND document_type='packing_list'",(customer_id,)).fetchone()
    c.close()
    return {**dict(row),"mappings":mappings,"row_rule":dict(rr) if rr else {}}

def _save_mappings(c,cid,mappings):
    for dtype,fields in mappings.items():
        for field,value in fields.items():
            if isinstance(value,dict):
                primary=str(value.get("primary","")).strip()
                alternatives=_headers(primary,value.get("alternatives",[]))
            else:
                primary=str(value or "").strip(); alternatives=""
            if primary or alternatives:
                c.execute("INSERT INTO customer_mappings(customer_id,document_type,standard_field,primary_header,alternative_headers) VALUES(?,?,?,?,?)",(cid,dtype,field,primary,alternatives))

def _save_row_rule(c,cid,rr):
    c.execute("DELETE FROM row_rules WHERE customer_id=? AND document_type='packing_list'",(cid,))
    c.execute("INSERT INTO row_rules(customer_id,document_type,item_identifier,aggregation_labels,continue_after_aggregation) VALUES(?,?,?,?,?)",(cid,"packing_list",rr.get("item_identifier","ITEM NAME"),rr.get("aggregation_labels","TOTAL;SUBTOTAL;GRAND TOTAL"),1 if rr.get("continue_after_aggregation",True) else 0))

@router.post("")
def create_customer(data:CustomerIn):
    name=data.customer_name.strip(); normalized=norm(name)
    if not name: raise HTTPException(400,"Nama customer wajib diisi")
    c=get_conn()
    try:
        cur=c.execute("INSERT INTO customers(customer_name,normalized_name,status) VALUES(?,?,?)",(name,normalized,data.status)); cid=cur.lastrowid
        _save_mappings(c,cid,data.mappings); _save_row_rule(c,cid,data.row_rule)
        c.commit(); return get_customer(cid)
    except Exception as e:
        c.rollback(); raise HTTPException(400,str(e))
    finally: c.close()

@router.put("/id/{customer_id}")
def update_customer(customer_id:int,data:CustomerIn):
    c=get_conn()
    if not c.execute("SELECT 1 FROM customers WHERE id=?",(customer_id,)).fetchone(): c.close(); raise HTTPException(404,"Customer tidak ditemukan")
    try:
        c.execute("UPDATE customers SET customer_name=?,normalized_name=?,status=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(data.customer_name.strip(),norm(data.customer_name),data.status,customer_id))
        c.execute("DELETE FROM customer_mappings WHERE customer_id=?",(customer_id,)); _save_mappings(c,customer_id,data.mappings); _save_row_rule(c,customer_id,data.row_rule)
        c.commit()
    except Exception as e:
        c.rollback(); raise HTTPException(400,str(e))
    finally: c.close()
    return get_customer(customer_id)

@router.get("/export/excel")
def export_customers():
    c=get_conn(); rows=[dict(r) for r in c.execute("SELECT id,customer_name,status FROM customers ORDER BY customer_name")]
    maps=[]
    for r in c.execute("SELECT c.customer_name,m.document_type,m.standard_field,m.primary_header,m.alternative_headers FROM customer_mappings m JOIN customers c ON c.id=m.customer_id ORDER BY c.customer_name,m.document_type,m.standard_field"): maps.append(dict(r))
    rules=[]
    for r in c.execute("SELECT c.customer_name,r.item_identifier,r.aggregation_labels,r.continue_after_aggregation FROM row_rules r JOIN customers c ON c.id=r.customer_id WHERE r.document_type='packing_list' ORDER BY c.customer_name"): rules.append(dict(r))
    c.close()
    wb=Workbook(); ws=wb.active; ws.title="Customers"
    ws.append(["Customer Name","Status"])
    for r in rows: ws.append([r["customer_name"],r["status"]])
    wm=wb.create_sheet("Mappings"); wm.append(["Customer Name","Document Type","Standard Field","Primary Header","Alternative Headers"])
    for r in maps: wm.append([r["customer_name"],r["document_type"],r["standard_field"],r["primary_header"],r["alternative_headers"]])
    wr=wb.create_sheet("Row Rules"); wr.append(["Customer Name","Item Identifier","Aggregation Labels","Continue After Aggregation"])
    for r in rules: wr.append([r["customer_name"],r["item_identifier"],r["aggregation_labels"],"TRUE" if r["continue_after_aggregation"] else "FALSE"])
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=customer-master.xlsx"})

@router.post("/import/excel")
async def import_customers(file:UploadFile=File(...)):
    if not (file.filename or "").lower().endswith((".xlsx",".xlsm")):
        raise HTTPException(400,"File Customer Master harus .xlsx atau .xlsm")
    data=await file.read()
    try: wb=load_workbook(io.BytesIO(data),data_only=True)
    except Exception as e: raise HTTPException(400,f"Excel tidak dapat dibaca: {e}")
    if "Customers" not in wb.sheetnames or "Mappings" not in wb.sheetnames:
        raise HTTPException(400,"Template wajib memiliki sheet Customers dan Mappings")
    ws=wb["Customers"]; wm=wb["Mappings"]; wr=wb["Row Rules"] if "Row Rules" in wb.sheetnames else None
    def rows(ws):
        vals=list(ws.iter_rows(values_only=True));
        if not vals:return []
        headers=[str(x or "").strip() for x in vals[0]]
        return [dict(zip(headers,r)) for r in vals[1:] if any(x not in (None,"") for x in r)]
    cust_rows=rows(ws); map_rows=rows(wm); rule_rows=rows(wr) if wr else []
    c=get_conn(); imported=0
    try:
        for cr in cust_rows:
            name=str(cr.get("Customer Name") or "").strip()
            if not name: continue
            normalized=norm(name); status=str(cr.get("Status") or "ACTIVE").strip().upper()
            existing=c.execute("SELECT id FROM customers WHERE normalized_name=?",(normalized,)).fetchone()
            if existing: cid=existing[0]; c.execute("UPDATE customers SET customer_name=?,status=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(name,status,cid))
            else: cid=c.execute("INSERT INTO customers(customer_name,normalized_name,status) VALUES(?,?,?)",(name,normalized,status)).lastrowid
            c.execute("DELETE FROM customer_mappings WHERE customer_id=?",(cid,))
            for mr in map_rows:
                if norm(str(mr.get("Customer Name") or ""))!=normalized: continue
                dtype=str(mr.get("Document Type") or "").strip(); field=str(mr.get("Standard Field") or "").strip()
                if dtype and field:
                    c.execute("INSERT INTO customer_mappings(customer_id,document_type,standard_field,primary_header,alternative_headers) VALUES(?,?,?,?,?)",(cid,dtype,field,str(mr.get("Primary Header") or "").strip(),str(mr.get("Alternative Headers") or "").strip()))
            rr=next((x for x in rule_rows if norm(str(x.get("Customer Name") or ""))==normalized),None)
            _save_row_rule(c,cid,rr or {})
            imported+=1
        c.commit()
    except Exception as e:
        c.rollback(); raise HTTPException(400,f"Import gagal: {e}")
    finally: c.close()
    return {"status":"ok","imported":imported}
